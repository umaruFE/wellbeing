from pathlib import Path
from docx import Document


SOURCE = Path(r"D:\workspace\wellbeing\Wellbeing_AI教育内容平台功能清单_已实现功能修订版.docx")
OUTPUT = Path(r"D:\workspace\wellbeing\Wellbeing_AI教育内容平台功能清单_四模块已实现修订版.docx")


REPLACEMENTS = {
    "版本  V2026.09.20-R2     更新日期  2026 年 9 月 20 日":
        "版本  V2026.09.20-R3     更新日期  2026 年 9 月 20 日",
    "根据项目当前实现更新 AI 绘本制作、星光录音棚和歌曲编唱屋的功能点，统一模块名称与实际制作流程；能力训练、趣味练习、主题活动继续仅保留名称和待确认说明。":
        "根据项目当前实现更新 AI 绘本制作、星光录音棚、歌曲编唱屋和互动式情景瑜伽的功能点，统一模块名称与实际制作流程；能力训练、趣味练习、主题活动继续仅保留名称和待确认说明。",
    "4  互动式情景瑜伽": "4  互动式情景瑜伽（已实现）",
    "将英语教学情境、故事任务与儿童瑜伽动作结合，让教师通过逐页动态课件带领学生在理解和使用英语的同时参与身体活动。AI 从教学需求设定、情境与动作编排、教学方案生成到视觉和声音制作提供辅助，最终形成可直接用于课堂的沉浸式情景瑜伽内容。":
        "通过四步工作室把英语教学目标、主题情境和儿童瑜伽动作编排为可编辑的逐页课件，并以独立窗口直接进入课堂授课。",
    "教学需求设定": "第一步  基本信息",
    "教师输入课程主题及故事或情境描述，并选择适合的课堂场景":
        "教师选择学生年龄、活动时长和一个或多个情境主题",
    "设置学生年龄段或年级，并输入重点英语词汇、重点句型和学习目标":
        "设置可用道具、目标语言和特殊要求；目标语言为生成活动方案的必填依据",
    "标准动作库与 AI 编排": "第二步  活动方案",
    "平台建设标准儿童瑜伽动作库，统一管理体式名称、动作描述、难度和适配情境":
        "AI 根据年龄、时长、主题、道具和目标语言生成活动名称、故事内容、英语目标、幸福力目标、成果目标和材料准备",
    "AI 根据课程主题、场景、学生年龄和教学内容，从标准动作库中匹配适合的体式并编排动作序列":
        "活动方案各字段均可编辑，并支持按当前基本信息重新生成",
    "AI 自动生成完整的逐页教学情境，每页包含学生端英语文案、对应瑜伽动作和场景画面描述，并形成连贯的故事与身体活动节奏":
        "教师确认活动方案后进入页面设计，作品在各步骤持续保存，可从作品列表继续编辑",
    "AI 配套教学方案": "第三步  页面设计",
    "在生成学生端课件的同时，AI 同步生成配套教学方案，供教师备课和课堂引导使用":
        "AI 按活动方案生成逐页结构，支持场景、过渡、动作、回归和结尾等页面类型",
    "教学方案可包含逐页教师引导语、表情与语气、动作示范、课堂互动建议、材料与道具提示，以及常见课堂情况的应对建议":
        "每页可编辑页面类型、瑜伽体式、学生端英文文案、画面描述和教师引导语；动作页从标准儿童瑜伽体式库选择动作并显示动作说明",
    "内容审核与调整": "页面编辑与调整",
    "教师可通过 AI 对整套内容提出修改要求，也可针对指定单页进行调整":
        "支持新增、删除页面，并调整页面内容和动作编排",
    "支持针对英语文案、瑜伽动作和场景描述等内容进行调整；内容确认后进入画面生成阶段":
        "支持重新生成整套页面设计，确认后进入插图制作",
    "如后续需要修改已确认内容，可返回对应页面调整，并重新生成受影响的画面或动态内容":
        "已保存作品可返回任一步骤继续修改；页面变更后可重新制作受影响的插图",
    "画面生成与动态化": "第四步  制作插图",
    "AI 根据每页确认的内容生成对应场景画面，生成结果支持修改或重新生成":
        "AI 根据每页确认的场景、体式和画面描述生成 16:9 无文字插图",
    "通过机构 IP 角色动态展示瑜伽动作从起始到结束的完整过程，便于教师带领学生跟随练习":
        "支持整套批量生成，也可对指定单页重新生成",
    "根据场景需要加入流水、飞鸟等局部环境微动画，增强情境的沉浸感，同时避免干扰核心动作展示":
        "生成状态和失败信息按页展示，已完成图片与页面内容一并保存",
    "AI 语音、音乐与场景音效": "作品管理",
    "AI 根据每页学生端英语内容自动生成英语语音，用于课堂播放和语言输入；不提供教师上传音频替换该语音的功能":
        "作品列表支持搜索、继续编辑和删除，并显示页面插图完成进度",
    "AI 自动生成基础版本的背景音乐，并根据不同页面的情境生成匹配的环境音效":
        "已生成页面的作品可从列表直接进入授课，也可在制作步骤中进入授课模式",
    "教师可对背景音乐和环境音效进行调整或替换，也可上传自有音乐和音效进行配置":
        "中英文界面均已适配，切换界面语言不改变已保存的业务数据",
    "课堂呈现": "课堂授课",
    "最终形成类似逐页课件的动态情景瑜伽内容，教师可根据学生状态自主控制翻页和课堂节奏":
        "授课模式以独立窗口展示逐页情景瑜伽课件，教师可自主控制翻页和课堂节奏",
    "课堂大屏主要展示场景画面、学生英语内容、IP 角色动作及相关多媒体效果；教师端可查看配套教学方案进行引导":
        "课件展示页面插图和学生端英文文案，支持键盘前后翻页及 Esc 退出",
    "教师带领学生完成情境中的瑜伽动作和英语互动，支持大屏、触摸屏及与其他教学素材组合使用":
        "教师结合页面中的瑜伽体式和引导内容，带领学生完成英语表达与身体活动",
    "客户价值  将课程主题与英语学习目标转化为可审核、可调整、可直接授课的情景瑜伽课件和配套教学方案，减少教师从零设计的工作。学生通过情境理解、英语输入和身体参与完成语言学习。":
        "客户价值  将英语学习目标与儿童瑜伽动作转化为可编辑、可生成插图、可直接授课的逐页活动，减少教师从零设计和制作课件的工作。",
}


def replace_visible_text(paragraph, new_text):
    nodes = paragraph._p.xpath('.//w:t')
    if not nodes:
        paragraph.add_run(new_text)
        return
    nodes[0].text = new_text
    for node in nodes[1:]:
        node.text = ''


def replace_fragment(paragraph, old_text, new_text):
    for node in paragraph._p.xpath('.//w:t'):
        if old_text in (node.text or ''):
            node.text = node.text.replace(old_text, new_text)
            return True
    return False


def iter_story_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for story in (section.header, section.footer):
            yield from story.paragraphs
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def main():
    doc = Document(SOURCE)
    replaced = set()
    for paragraph in iter_story_paragraphs(doc):
        old = paragraph.text
        if old in REPLACEMENTS:
            replace_visible_text(paragraph, REPLACEMENTS[old])
            replaced.add(old)
        replace_fragment(paragraph, "V2026.09.20-R2  |  2026-09-20", "V2026.09.20-R3  |  2026-09-20")

    missing = sorted(set(REPLACEMENTS) - replaced)
    if missing:
        raise RuntimeError("未找到以下待替换文本：\n" + "\n".join(missing))

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
